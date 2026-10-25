import json
import os
import re
import unittest
from datetime import datetime

from django.conf import settings
from django.core.management.base import BaseCommand as DjangoBaseCommand
from django.test.runner import DiscoverRunner
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from aplicaciones.casos_prueba import CasoPruebaMixin

REQUISITOS = {
    'RF-001': 'Registro de empleados', 'RF-002': 'Registro de horas extra',
    'RF-003': 'Solicitud de certificados laborales', 'RF-004': 'Consulta de nómina, reportes y PDF',
    'RF-005': 'Bonificaciones configurables', 'RF-006': 'Frecuencia y método de pago',
    'RF-007': 'Aprobación de solicitudes', 'RF-008': 'Cálculo de nómina y pre-nómina',
    'RF-009': 'Registro de usuarios y contraseña cifrada', 'RF-010': 'Autenticación con correo y contraseña',
    'RF-011': 'Cambio de contraseña', 'RF-012': 'Perfil de usuario', 'RF-013': 'Roles y permisos diferenciados',
    'RF-014': 'Administración de roles', 'RF-015': 'Asignación de roles', 'RF-016': 'Consulta de usuarios y roles con filtros',
    'RF-017': 'Bitácora de auditoría', 'RNF-001': 'Cifrado de contraseñas', 'RNF-003': 'Control de acceso por roles',
    'RNF-004': 'Complejidad de clave', 'RNF-005': 'Autenticación <= 2 s',
    'RNF-006': 'Nómina de 200 empleados <= 1 min', 'RNF-007': 'Consultas filtradas <= 3 s',
    'RNF-009': 'Mensajes de error claros', 'RNF-011': 'Transaccionalidad', 'RNF-016': 'Volumen de 200 empleados',
}


def _orden(rf):
    m = re.search(r'(\d+)', rf)
    return (1 if rf.startswith('RNF') else 0, int(m.group(1)) if m else 999)


ENCABEZADO = ['CP', 'RF', 'Descripción', 'Input', 'Resultado esperado', 'Resultado obtenido', 'Observación', 'Estado']


def _estado(c):
    if c.get('estado') != 'ok':
        return 'No aprobado'
    return 'Observación' if c.get('observacion_estado') else 'Aprobado'


def _fila(c):
    nota = c.get('observacion', '')
    if c.get('id', '').startswith('aplicaciones.'):
        nota = nota + ('; ' if nota else '') + 'Prueba automatizada: ' + '.'.join(c['id'].split('.')[-2:])
    elif c.get('id'):
        nota = nota + ('; ' if nota else '') + 'Prueba de carga: ' + c['id']
    return [c.get('cp'), c.get('rf'), c.get('descripcion'), c.get('input'), c.get('esperado'),
            c.get('obtenido'), nota, _estado(c)]


class _Resultado(unittest.TextTestResult):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.registros = []

    def _agregar(self, test, estado, error=''):
        metodo = getattr(test, getattr(test, '_testMethodName', ''), None)
        if not getattr(metodo, '_cp_rf', None):
            return
        ident = test.id()
        real = CasoPruebaMixin.resultados().get(ident)
        if error:
            error = re.sub(r'\s+', ' ', error)[:300]
        self.registros.append({'id': ident, 'rf': metodo._cp_rf, 'descripcion': metodo._cp_descripcion,
            'input': metodo._cp_entrada, 'esperado': metodo._cp_esperado,
            'obtenido': real or (error if estado != 'ok' else 'Se obtuvo el resultado esperado.'),
            'estado': estado, 'error': error, 'observacion': metodo._cp_observacion or '',
            'observacion_estado': metodo._cp_observacion_estado})

    def addSuccess(self, test): super().addSuccess(test); self._agregar(test, 'ok')
    def addFailure(self, test, err): super().addFailure(test, err); self._agregar(test, 'fallo', self._exc_info_to_string(err, test))
    def addError(self, test, err): super().addError(test, err); self._agregar(test, 'error', self._exc_info_to_string(err, test))


class _Runner(DiscoverRunner):
    def get_resultclass(self): return _Resultado


class Command(DjangoBaseCommand):
    help = 'Ejecuta y documenta los casos de prueba por requisito funcional.'

    def add_arguments(self, parser):
        parser.add_argument('--salida', default='docs/Plan_de_Pruebas.xlsx')
        parser.add_argument('--solo-generar', action='store_true')

    def handle(self, *args, **opts):
        CasoPruebaMixin.limpiar_resultados()
        resultado_carga = []
        if not opts['solo_generar']:
            runner = _Runner(verbosity=1, interactive=False)
            runner.setup_test_environment()
            suite = runner.build_suite(['aplicaciones'])
            bases = runner.setup_databases()
            try:
                result = runner.run_suite(suite)
            finally:
                runner.teardown_databases(bases)
                runner.teardown_test_environment()
            resultado_carga = result.registros
        else:
            previo = os.path.join(settings.BASE_DIR, 'docs', 'resultados_pruebas.json')
            if os.path.exists(previo):
                with open(previo, encoding='utf-8') as f:
                    resultado_carga = json.load(f).get('casos', [])
        destino_docs = os.path.join(settings.BASE_DIR, 'docs')
        os.makedirs(destino_docs, exist_ok=True)
        rendimiento = os.path.join(destino_docs, 'resultados_rendimiento.json')
        if os.path.exists(rendimiento):
            with open(rendimiento, encoding='utf-8') as f:
                datos_rendimiento = json.load(f)
            resultado_carga.extend(datos_rendimiento.get('casos', []))
        resultado_carga.sort(key=lambda c: (_orden(c.get('rf', 'RF-999')), c.get('id', '')))
        for i, c in enumerate(resultado_carga, 1): c['cp'] = 'CP{0}'.format(i)
        fecha = datetime.now().astimezone().isoformat(timespec='seconds')
        aprobados = sum(c.get('estado') == 'ok' for c in resultado_carga)
        ruta_json = os.path.join(destino_docs, 'resultados_pruebas.json')
        with open(ruta_json, 'w', encoding='utf-8') as f:
            json.dump({'fecha_ejecucion': fecha, 'totales': {'casos': len(resultado_carga), 'aprobados': aprobados,
                'fallidos': sum(c.get('estado') in ('fallo','error') for c in resultado_carga)}, 'casos': resultado_carga}, f, ensure_ascii=False, indent=2)
        self._excel(os.path.join(settings.BASE_DIR, opts['salida']), resultado_carga)
        self._markdown(os.path.join(destino_docs, 'Plan_de_Pruebas.md'), resultado_carga, fecha)
        self.stdout.write('Plan generado: {0}; resultados: {1}'.format(opts['salida'], ruta_json))

    def _excel(self, ruta, casos):
        wb = Workbook(); plan = wb.active; plan.title = 'Plan de pruebas'; req = wb.create_sheet('Requisitos')
        plan.append(ENCABEZADO)
        for c in casos:
            plan.append(_fila(c))
        req.append(['RF/RNF','Descripción','Cantidad de CP','Aprobados'])
        for rf, desc in REQUISITOS.items():
            fila = req.max_row + 1
            req.append([rf,desc,'=COUNTIF(\'Plan de pruebas\'!$B:$B,A{0})'.format(fila),'=COUNTIFS(\'Plan de pruebas\'!$B:$B,A{0},\'Plan de pruebas\'!$H:$H,"Aprobado")'.format(fila)])
        navy = '1E1B4B'; blanco = 'FFFFFF'
        for ws in (req, plan):
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]: cell.fill=PatternFill('solid', fgColor=navy); cell.font=Font(name='Arial', bold=True, color=blanco)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font=Font(name='Arial', color='000000'); cell.alignment=Alignment(vertical='top', wrap_text=True)
                    if cell.row % 2 == 0: cell.fill=PatternFill('solid', fgColor='F1F5F9')
            for col in range(1, ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width = [14,22,42,45,45,45,45,16][col-1]
        for estado,color in [('Aprobado','DCFCE7'),('No aprobado','FEE2E2'),('Observación','FEF3C7')]:
            plan.conditional_formatting.add('H2:H{0}'.format(plan.max_row), FormulaRule(formula=['H2="{0}"'.format(estado)], fill=PatternFill('solid', fgColor=color)))
        wb.calculation.fullCalcOnLoad = True
        os.makedirs(os.path.dirname(ruta), exist_ok=True); wb.save(ruta)

    def _markdown(self, ruta, casos, fecha):
        def celda(valor):
            return str(valor or '').replace('|', '\\|').replace('\n', ' ')
        estados = [_estado(c) for c in casos]
        lineas = ['# Plan de pruebas', '',
                  'Ejecución: {0}. Casos: {1}; aprobados: {2}; no aprobados: {3}; con observación: {4}.'.format(
                      fecha, len(casos), estados.count('Aprobado'), estados.count('No aprobado'), estados.count('Observación')),
                  '', '| ' + ' | '.join(ENCABEZADO) + ' |', '|' + '---|' * len(ENCABEZADO)]
        for c in casos:
            lineas.append('| ' + ' | '.join(celda(v) for v in _fila(c)) + ' |')
        with open(ruta, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lineas) + '\n')
